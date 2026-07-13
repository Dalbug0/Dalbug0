from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt


OUTPUT_PATH = Path(__file__).with_name("Заявление_на_дивиденды.docx")
FONT_NAME = "Times New Roman"


def set_run_font(run, size=Pt(14), bold=False, italic=False):
    run.font.name = FONT_NAME
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    return run


def remove_table_borders(table):
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "nil")


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)

    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_paragraph_format(paragraph, *, alignment=None, before=0, after=0, line_spacing=True):
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    if line_spacing:
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return paragraph


def add_text_paragraph(container, text, **format_options):
    paragraph = container.add_paragraph()
    set_paragraph_format(paragraph, **format_options)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_labeled_line(container, label="", *, width=Inches(3.62), after=6):
    paragraph = container.add_paragraph()
    set_paragraph_format(paragraph, after=after)
    if label:
        set_run_font(paragraph.add_run(label))
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        width,
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.LINES,
    )
    set_run_font(paragraph.add_run("\t"))
    return paragraph


def add_full_line(container, *, width, after=0):
    paragraph = container.add_paragraph()
    set_paragraph_format(paragraph, after=after)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        width,
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.LINES,
    )
    set_run_font(paragraph.add_run("\t"))
    return paragraph


def add_caption(container, text, *, after=8):
    paragraph = container.add_paragraph()
    set_paragraph_format(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=after,
        line_spacing=False,
    )
    set_run_font(paragraph.add_run(text), size=Pt(10))
    return paragraph


def build_document():
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal_style = document.styles["Normal"]
    normal_style.font.name = FONT_NAME
    normal_style.font.size = Pt(14)
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal_style.paragraph_format.space_after = Pt(0)

    document.core_properties.title = "Заявление на перечисление дивидендов"
    document.core_properties.subject = "Заявление акционера"

    usable_width = section.page_width - section.left_margin - section.right_margin
    header_table = document.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = int(usable_width * 0.40)
    header_table.columns[1].width = int(usable_width * 0.60)
    remove_table_borders(header_table)

    left_cell, right_cell = header_table.rows[0].cells
    left_cell.width = int(usable_width * 0.40)
    right_cell.width = int(usable_width * 0.60)
    right_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(left_cell)
    set_cell_margins(right_cell)

    first_paragraph = right_cell.paragraphs[0]
    set_paragraph_format(first_paragraph, after=6)
    set_run_font(
        first_paragraph.add_run(
            "Главному бухгалтеру ОАО «Белэнергоремналадка»"
        )
    )
    add_labeled_line(right_cell, "Ф.И.О. акционера ")
    add_full_line(right_cell, width=Inches(3.62), after=6)
    add_labeled_line(right_cell, "Адрес ")
    add_full_line(right_cell, width=Inches(3.62), after=6)
    add_labeled_line(right_cell, "Контактный телефон ")
    add_labeled_line(right_cell, "Паспорт номер ")
    add_labeled_line(right_cell, "Кем выдан ")
    add_labeled_line(right_cell, "Дата выдачи ")
    add_labeled_line(right_cell, "Идентификационный № ", after=0)

    title = document.add_paragraph()
    set_paragraph_format(
        title,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        before=30,
        after=22,
    )
    title_run = set_run_font(title.add_run("ЗАЯВЛЕНИЕ"), bold=True)
    title_run._element.get_or_add_rPr().append(OxmlElement("w:spacing"))
    title_run._element.rPr[-1].set(qn("w:val"), "30")

    request = document.add_paragraph()
    set_paragraph_format(
        request,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        after=0,
    )
    set_run_font(
        request.add_run(
            "Прошу перечислять причитающиеся мне дивиденды на банковский счет"
        )
    )
    add_full_line(document, width=usable_width)
    add_caption(document, "(номер счета)")

    bank = document.add_paragraph()
    set_paragraph_format(
        bank,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        after=0,
    )
    set_run_font(bank.add_run("в банке "))
    bank.paragraph_format.tab_stops.add_tab_stop(
        usable_width,
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.LINES,
    )
    set_run_font(bank.add_run("\t"))
    add_caption(document, "(наименование банка)", after=12)

    consent = document.add_paragraph()
    set_paragraph_format(
        consent,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        before=10,
        after=26,
    )
    set_run_font(
        consent.add_run(
            "Настоящим подтверждаю свое согласие несогласие "
            "(НУЖНО ПОДЧЕРКНУТЬ) на внесение, обработку, хранение и передачу "
            "моих персональных данных, указанных в заявлении."
        ),
        italic=True,
    )

    footer_table = document.add_table(rows=1, cols=2)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    footer_table.autofit = False
    footer_table.columns[0].width = int(usable_width * 0.50)
    footer_table.columns[1].width = int(usable_width * 0.50)
    remove_table_borders(footer_table)

    signature_cell, date_cell = footer_table.rows[0].cells
    signature_cell.width = int(usable_width * 0.50)
    date_cell.width = int(usable_width * 0.50)
    set_cell_margins(signature_cell, end=120)
    set_cell_margins(date_cell, start=120)

    signature = signature_cell.paragraphs[0]
    set_paragraph_format(signature, before=12)
    set_run_font(signature.add_run("Подпись "))
    signature.paragraph_format.tab_stops.add_tab_stop(
        Inches(2.10),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.LINES,
    )
    set_run_font(signature.add_run("\t"))

    date = date_cell.paragraphs[0]
    set_paragraph_format(date, before=12)
    set_run_font(date.add_run("Дата "))
    date.paragraph_format.tab_stops.add_tab_stop(
        Inches(2.45),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.LINES,
    )
    set_run_font(date.add_run("\t"))

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(f"Создан файл: {output}")
